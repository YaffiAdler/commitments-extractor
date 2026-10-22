from openpyxl import load_workbook

from app.export.xlsx import (
    COLUMNS,
    EXTRA_HEADERS,
    IRRELEVANT_SHEET_TITLE,
    NEEDS_REVIEW_YES,
    REVIEW_FILL,
    SPLIT_UNCERTAIN_TEXT,
    write_xlsx,
)
from app.schemas import SPLIT_UNCERTAIN, Commitment, Field, FormGroup, IrrelevantPage, Job


def _job(commitment: Commitment) -> Job:
    return Job(groups=[FormGroup(source_file="a.pdf", pages=[0], commitment=commitment)])


def test_writes_header_and_row(tmp_path):
    commitment = Commitment(
        id_number=Field[str](value="012345678", confidence=1.0, needs_review=False),
        central_numbers=[Field[str](value="111111"), Field[str](value="222222")],
    )
    sheet = load_workbook(write_xlsx(_job(commitment), tmp_path / "out.xlsx")).active
    assert sheet.sheet_view.rightToLeft
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == tuple(h for h, _ in COLUMNS) + EXTRA_HEADERS
    assert rows[1][2] == "012345678"  # leading zero kept
    assert rows[1][3] == "111111; 222222"
    assert rows[1][len(COLUMNS):] == ("a.pdf", "1", None, NEEDS_REVIEW_YES)
    assert workbook_sheets(tmp_path / "out.xlsx") == [sheet.title]


def workbook_sheets(path):
    return load_workbook(path).sheetnames


def test_uncertain_split_and_irrelevant_pages(tmp_path):
    group = FormGroup(source_file="a.pdf", pages=[1, 2], split_confidence=0.3, split_reasons=[SPLIT_UNCERTAIN])
    job = Job(groups=[group], irrelevant_pages=[IrrelevantPage(source_file="a.pdf", page=0)])
    workbook = load_workbook(write_xlsx(job, tmp_path / "out.xlsx"))
    row = list(workbook.active.iter_rows(min_row=2, values_only=True))[0]
    assert row[len(COLUMNS):] == ("a.pdf", "2, 3", SPLIT_UNCERTAIN_TEXT, NEEDS_REVIEW_YES)
    assert list(workbook[IRRELEVANT_SHEET_TITLE].iter_rows(min_row=2, values_only=True)) == [("a.pdf", 1)]


def test_missing_fields_are_empty_and_highlighted(tmp_path):
    sheet = load_workbook(write_xlsx(_job(Commitment()), tmp_path / "out.xlsx")).active
    for column in range(1, len(COLUMNS) + 1):
        cell = sheet.cell(row=2, column=column)
        assert cell.value is None
        assert cell.fill.start_color.rgb.endswith(REVIEW_FILL.start_color.rgb[-6:])


def test_optional_empty_list_is_not_highlighted(tmp_path):
    commitment = Commitment(optional_fields=["commitment_numbers"])
    sheet = load_workbook(write_xlsx(_job(commitment), tmp_path / "out.xlsx")).active
    column = [attribute for _, attribute in COLUMNS].index("commitment_numbers") + 1
    assert sheet.cell(row=2, column=column).fill.fill_type is None
    assert sheet.cell(row=2, column=column - 1).fill.fill_type == REVIEW_FILL.fill_type
