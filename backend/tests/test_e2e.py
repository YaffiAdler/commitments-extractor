"""End to end: upload a PDF -> one row in the job -> a valid xlsx, using the fake provider."""
from io import BytesIO

import pymupdf
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings
from app.config import ExportSettings
from app.export.xlsx import HEADERS
from app.main import create_app
from app.providers import fake
from app.providers.fake import FakeProvider


def _one_page_pdf() -> bytes:
    with pymupdf.open() as doc:
        doc.new_page().insert_text((72, 72), "dummy form")
        return doc.tobytes()


def test_upload_to_xlsx(tmp_path, monkeypatch):
    monkeypatch.setenv("CE_TMP_DIR", str(tmp_path))
    monkeypatch.setenv("CE_RENDER_DPI", "72")
    client = TestClient(create_app(Settings(), FakeProvider()))

    created = client.post("/jobs", files={"files": ("form.pdf", _one_page_pdf(), "application/pdf")})
    assert created.status_code == 201
    job_id = created.json()["id"]

    job = client.get(f"/jobs/{job_id}").json()
    assert job["status"] == "done"
    assert len(job["groups"]) == 1
    assert job["groups"][0]["commitment"]["id_number"]["value"] == fake.FAKE_ID_NUMBER

    exported = client.post(f"/jobs/{job_id}/export")
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == tuple(HEADERS[key] for key in ExportSettings().columns)
    assert len(rows) == 2
    assert rows[1][2] == fake.FAKE_ID_NUMBER
    assert rows[1][6] == fake.FAKE_COMMITMENT_NUMBER

    assert client.delete(f"/jobs/{job_id}").status_code == 204
    assert client.get(f"/jobs/{job_id}").status_code == 404
    assert not (tmp_path / job_id).exists()
