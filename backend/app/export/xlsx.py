"""Export a finished job to an Excel workbook, one row per form.

The columns and their order come from ``config.ExportSettings``. Fields that
need review get an orange background, and the "הערות לבדיקה" column sums up
why (field reasons and form-level reasons as Hebrew text). Empty fields stay
empty. Numbers are text cells (leading zeros matter); the date is a real date
cell. Pages that belong to no form are listed on a second sheet.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.config import ExportSettings, MultiValueMode
from app.pipeline.reason_texts import FORM_REASON_TEXTS, REASON_TEXTS, explain
from app.schemas import COMMITMENT_FIELDS, NOT_FOUND, SPLIT_UNCERTAIN, Commitment, Field, FormGroup, Job

SHEET_TITLE = "התחייבויות"
IRRELEVANT_SHEET_TITLE = "עמודים לא רלוונטיים"

REVIEW_NOTES = "review_notes"
SOURCE_FILE = "source_file"
PAGES = "pages"

# Header text per column key (``config.EXPORT_COLUMN_KEYS``); exported verbatim.
HEADERS: dict[str, str] = {
    "first_name": "שם פרטי",
    "last_name": "שם משפחה",
    "id_number": "ת.ז.",
    "central_numbers": "מספר מרכזי",
    "received_date": "תאריך קבלה",
    "sender": "גורם שולח",
    "commitment_numbers": "מספר התחייבות",
    REVIEW_NOTES: "הערות לבדיקה",
    SOURCE_FILE: "קובץ מקור",
    PAGES: "עמודים",
}
IRRELEVANT_HEADERS = (HEADERS[SOURCE_FILE], "עמוד")

LIST_FIELDS = frozenset({"central_numbers", "commitment_numbers"})
DATE_FIELDS = frozenset({"received_date"})
TEXT_FORMAT = "@"
DATE_FORMAT = "DD/MM/YYYY"
LIST_SEPARATOR = "; "
PAGE_SEPARATOR = ", "
NOTE_LINE_SEPARATOR = "\n"
NOTE_LABEL_SEPARATOR = ": "
# A field flagged without any reason code (not produced by the pipeline, but the schema allows it).
GENERIC_REVIEW_TEXT = "סומן לבדיקה"

ORANGE = "F4B183"
REVIEW_FILL = PatternFill(fill_type="solid", start_color=ORANGE, end_color=ORANGE)
HEADER_FONT = Font(bold=True)
NOTES_ALIGNMENT = Alignment(wrap_text=True, vertical="top")
FROZEN_CELL = "A2"
# Column width in characters: longest line + padding, kept within these bounds.
MIN_COLUMN_WIDTH = 8
MAX_COLUMN_WIDTH = 60
COLUMN_PADDING = 2


@dataclass(frozen=True)
class _Column:
    """One sheet column: a column key, and for a list field in ``columns``
    mode the index of the list item it shows."""

    key: str
    item: Optional[int] = None

    @property
    def header(self) -> str:
        base = HEADERS[self.key]
        return base if self.item is None else f"{base} {self.item + 1}"


def plan_columns(job: Job, settings: ExportSettings) -> list[_Column]:
    """The sheet's columns. In ``columns`` mode a list field takes as many
    columns as its longest list in the job (at least one)."""
    columns: list[_Column] = []
    for key in settings.columns:
        if key in LIST_FIELDS and settings.multi_values is MultiValueMode.COLUMNS:
            width = max((len(getattr(group.commitment, key)) for group in job.groups), default=0)
            columns.extend(_Column(key, item) for item in range(max(width, 1)))
        else:
            columns.append(_Column(key))
    return columns


def _joined(fields: list[Field[str]]) -> Optional[str]:
    return LIST_SEPARATOR.join(f.value for f in fields if f.value is not None) or None


def _field_cell(commitment: Commitment, column: _Column) -> tuple[str | date | None, bool]:
    """Value of a field column and whether it is highlighted."""
    value = getattr(commitment, column.key)
    if column.key not in LIST_FIELDS:
        return value.value, value.needs_review
    if column.item is None:
        return _joined(value), commitment.field_needs_review(column.key)
    if column.item < len(value):
        item = value[column.item]
        return item.value, item.needs_review
    # An empty list that must not be empty is flagged in its first column only.
    return None, column.item == 0 and commitment.field_needs_review(column.key)


def _texts(field: Field) -> list[str]:
    return field.explanations or explain(field.reasons) or [GENERIC_REVIEW_TEXT]


def _note(label: str, texts: list[str]) -> str:
    return label + NOTE_LABEL_SEPARATOR + LIST_SEPARATOR.join(texts)


def _field_notes(commitment: Commitment, name: str) -> list[str]:
    header = HEADERS[name]
    value = getattr(commitment, name)
    if name not in LIST_FIELDS:
        return [_note(header, _texts(value))] if value.needs_review else []
    if not value:
        return [_note(header, [REASON_TEXTS[NOT_FOUND]])] if commitment.field_needs_review(name) else []
    numbered = len(value) > 1
    return [_note(f"{header} {index + 1}" if numbered else header, _texts(item))
            for index, item in enumerate(value) if item.needs_review]


def review_notes(group: FormGroup) -> Optional[str]:
    """Why the form needs review, one line per form-level reason and per
    flagged field, in Hebrew. None if nothing is flagged. Describes the
    checks only, never the values."""
    form_codes = [SPLIT_UNCERTAIN] if SPLIT_UNCERTAIN in group.split_reasons else []
    form_codes += [code for code in group.form_reasons if code in FORM_REASON_TEXTS]
    lines = [FORM_REASON_TEXTS[code] for code in form_codes]
    for name in COMMITMENT_FIELDS:
        lines.extend(_field_notes(group.commitment, name))
    return NOTE_LINE_SEPARATOR.join(lines) or None


def _page_numbers(pages: list[int]) -> str:
    """1-based, as a reader counts pages."""
    return PAGE_SEPARATOR.join(str(page + 1) for page in pages)


def _format_field_cell(cell: Cell, key: str) -> None:
    """Set also on empty cells, so a value typed in by the reviewer keeps its
    leading zeros (text) or is read as a date."""
    cell.number_format = DATE_FORMAT if key in DATE_FIELDS else TEXT_FORMAT


def _write_row(sheet: Worksheet, row: int, group: FormGroup, columns: list[_Column]) -> None:
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=index)
        if column.key == REVIEW_NOTES:
            cell.value = review_notes(group)
            cell.alignment = NOTES_ALIGNMENT
        elif column.key == SOURCE_FILE:
            cell.value = group.source_file
        elif column.key == PAGES:
            cell.value = _page_numbers(group.pages)
        else:
            cell.value, highlighted = _field_cell(group.commitment, column)
            _format_field_cell(cell, column.key)
            if highlighted:
                cell.fill = REVIEW_FILL


def _display_length(value: object) -> int:
    if isinstance(value, date):
        return len(DATE_FORMAT)
    return max((len(line) for line in str(value).split(NOTE_LINE_SEPARATOR)), default=0)


def _fit_columns(sheet: Worksheet) -> None:
    """Width of each column from its longest line (openpyxl cannot measure text)."""
    for cells in sheet.iter_cols():
        longest = max((_display_length(cell.value) for cell in cells if cell.value is not None), default=0)
        width = min(max(longest + COLUMN_PADDING, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)
        sheet.column_dimensions[get_column_letter(cells[0].column)].width = width


def _start_sheet(sheet: Worksheet, headers: list[str]) -> None:
    sheet.sheet_view.rightToLeft = True
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
    sheet.freeze_panes = FROZEN_CELL


def _write_irrelevant_pages(sheet: Worksheet, job: Job) -> None:
    _start_sheet(sheet, list(IRRELEVANT_HEADERS))
    for page in job.irrelevant_pages:
        sheet.append([page.source_file, page.page + 1])
    _fit_columns(sheet)


def write_xlsx(job: Job, out_path: Path, settings: ExportSettings | None = None) -> Path:
    settings = settings or ExportSettings()
    columns = plan_columns(job, settings)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    _start_sheet(sheet, [column.header for column in columns])
    for row, group in enumerate(job.groups, start=2):
        _write_row(sheet, row, group, columns)
    _fit_columns(sheet)
    if job.irrelevant_pages:
        _write_irrelevant_pages(workbook.create_sheet(IRRELEVANT_SHEET_TITLE), job)
    workbook.save(out_path)
    return out_path
